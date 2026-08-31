from io import BytesIO

from docx import Document

from app.models import User, UserProfile
from app.schemas.tailored_resume import TailoredResumeContent
from app.services.resume_tailoring import (
    build_tailored_content,
    find_master_resume_filename,
    render_docx,
    validate_tailored_content,
)


def make_profile() -> UserProfile:
    return UserProfile(
        first_name="测试用户",
        email="candidate@example.com",
        phone="13800138000",
        portfolio_url="https://example.com",
        graduation_year="2027",
        highest_degree="硕士",
        work_history=[
            {
                "company": "AI公司",
                "title": "Agent实习生",
                "start_date": "2026-05",
                "end_date": "至今",
                "description": "负责Agent评测、Trace分析与工具编排。",
                "highlights": ["构建评测系统，准确率达到86%。"],
            },
            {
                "company": "工程公司",
                "title": "算法实习生",
                "start_date": "2026-01",
                "end_date": "2026-05",
                "description": "构建RAG知识库与FastAPI服务。",
                "highlights": ["处理3000+页文档，完成500+测试用例。"],
            },
        ],
        projects=[
            {
                "name": "Agent评测项目",
                "role": "核心开发",
                "start_date": "2026-05",
                "end_date": "至今",
                "description": "Agent评测与工具编排。",
                "technologies": ["Agent", "Python"],
                "highlights": ["人工核对准确率达到86%。"],
            },
            {
                "name": "地震模型",
                "role": "项目负责人",
                "start_date": "2025-07",
                "end_date": "2025-08",
                "description": "PyTorch与LoRA训练。",
                "technologies": ["PyTorch", "LoRA"],
                "highlights": ["实现预训练和微调流程。"],
            },
        ],
        education=[
            {
                "institution": "测试大学",
                "degree": "硕士",
                "field_of_study": "人工智能",
                "start_date": "2024-09",
                "end_date": "2027-06",
            }
        ],
        skills=["Python", "Agent", "RAG", "FastAPI", "Docker"],
        certificates=["GitHub开源项目2000+ Stars"],
        languages=["中文", "英语（CET-4）"],
        resume_original_filename="普通简历.pdf",
        resume_versions=[
            {
                "original_filename": "测试用户_完整母版.pdf",
                "file_path": "uploads/master.pdf",
                "uploaded_at": "2026-08-30T00:00:00+00:00",
                "is_current": True,
            }
        ],
    )


def test_build_tailored_content_prioritizes_job_evidence():
    profile = make_profile()
    jd = "负责Agent架构、工具编排、RAG知识库、自动化评测和高并发稳定性。"

    analysis, content, validation = build_tailored_content("AI研发工程师", jd, profile)

    assert content.target_role == "AI研发工程师"
    assert content.experiences[0].heading == "AI公司"
    assert content.core_skills[:3] == ["Agent", "RAG", "Python"]
    assert analysis["strong_count"] >= 2
    assert any(item["key"] == "performance" for item in analysis["requirements"])
    assert validation["valid"] is True
    assert validation["warnings"] == ["性能与稳定性相关要求缺少直接履历证据"]


def test_validation_rejects_numbers_not_found_in_master_profile():
    profile = make_profile()
    _, content, _ = build_tailored_content("AI研发工程师", "Agent与RAG", profile)
    changed = TailoredResumeContent.model_validate(content.model_dump())
    changed.experiences[0].bullets = ["将系统性能提升999倍。"]

    validation = validate_tailored_content(profile, changed)

    assert validation["valid"] is False
    assert "999" in validation["errors"][0]


def test_summary_uses_profile_facts_instead_of_fixed_ai_claims():
    profile = make_profile()
    profile.highest_degree = "本科"
    profile.education[0]["degree"] = "本科"
    profile.skills = ["Java", "SQL"]
    profile.work_history = [
        {
            "company": "软件公司",
            "title": "后端实习生",
            "description": "负责Java接口开发。",
            "highlights": ["完成接口开发与测试。"],
        }
    ]
    profile.projects = [
        {
            "name": "管理系统",
            "role": "开发",
            "description": "使用Java和SQL开发。",
            "highlights": ["完成管理功能。"],
        }
    ]

    _, content, validation = build_tailored_content(
        "后端开发工程师", "负责Java服务与SQL数据库开发。", profile
    )

    assert "本科" in content.summary
    assert "硕士" not in content.summary
    assert "Agent" not in content.summary
    assert "模型微调" not in content.summary
    assert validation["valid"] is True


def test_render_docx_contains_confirmed_profile_and_resume_content():
    profile = make_profile()
    user = User(email="candidate@example.com", password_hash="x", city="北京")
    _, content, _ = build_tailored_content("AI研发工程师", "Agent与RAG", profile)

    generated = render_docx(profile, user, content)
    document = Document(BytesIO(generated))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )

    assert generated.startswith(b"PK")
    assert "测试用户" in text
    assert "AI研发工程师" in text
    assert "Agent实习生" in text


def test_find_master_resume_prefers_current_version():
    assert find_master_resume_filename(make_profile()) == "测试用户_完整母版.pdf"
