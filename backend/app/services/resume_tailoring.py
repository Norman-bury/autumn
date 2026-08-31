import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.models import User, UserProfile
from app.schemas.tailored_resume import (
    RequirementMatch,
    ResumeEducation,
    ResumeEntry,
    TailoredResumeContent,
)

CATEGORY_RULES = (
    (
        "business",
        "需求理解与业务归因",
        ("需求", "业务", "归因", "数据挖掘", "问题定义"),
        ("评测", "归因", "badcase", "标准", "业务"),
    ),
    (
        "agent",
        "Agent架构与工具编排",
        ("agent", "智能体", "工具编排", "记忆管理", "推理策略", "任务拆解"),
        ("agent", "工具", "note", "上下文", "trace", "workflow"),
    ),
    (
        "knowledge",
        "RAG与知识工程",
        ("rag", "知识库", "召回", "上下文注入", "api 接入"),
        ("rag", "知识库", "向量", "检索", "markdown", "fastapi"),
    ),
    (
        "evaluation",
        "自动评测与持续迭代",
        ("评测", "回测", "case", "观测", "追踪", "效果"),
        ("评测", "测试", "准确率", "trace", "badcase", "计分"),
    ),
    (
        "model",
        "模型训练与推理",
        ("sft", "rl", "微调", "vllm", "ollama", "cv", "nlp"),
        ("sft", "lora", "微调", "vllm", "pytorch", "计算机视觉"),
    ),
    (
        "engineering",
        "AI工程与编程",
        ("python", "java", "javascript", "工程", "sdk", "api服务", "coding"),
        ("python", "fastapi", "docker", "linux", "git", "ci/cd", "算法"),
    ),
    (
        "performance",
        "性能与稳定性",
        ("高并发", "异步", "降级", "稳定运行", "延迟优化", "kv cache"),
        ("高并发", "异步", "降级", "并行计算", "稳定性", "延迟"),
    ),
    (
        "influence",
        "开源与技术影响力",
        ("github", "开源", "技术博客", "社区影响力", "个人 demo"),
        ("github", "stars", "开源", "skill", "软著", "论文"),
    ),
)

TECH_KEYWORDS = (
    "Agent",
    "RAG",
    "Context Engineering",
    "Prompt Engineering",
    "LangChain",
    "LangGraph",
    "Python",
    "Java",
    "JavaScript",
    "FastAPI",
    "vLLM",
    "Ollama",
    "Docker",
    "Linux",
    "CI/CD",
    "SFT",
    "LoRA",
    "DPO",
    "PyTorch",
    "MCP",
    "CV",
    "NLP",
)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return " ".join(_text(item) for item in value)
    if isinstance(value, dict):
        return " ".join(_text(item) for item in value.values())
    return str(value)


def _contains(text: str, term: str) -> bool:
    return term.lower() in text.lower()


def _period(item: dict[str, Any]) -> str | None:
    start = str(item.get("start_date") or "").strip()
    end = str(item.get("end_date") or "").strip()
    if start and end:
        return f"{start} - {end}"
    return start or end or None


def _source_items(profile: UserProfile) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for index, item in enumerate(profile.work_history or []):
        items.append(
            {
                "id": f"work-{index}",
                "kind": "work",
                "label": f"{item.get('company', '')} · {item.get('title', '')}".strip(
                    " ·"
                ),
                "data": item,
                "text": _text(item),
            }
        )
    for index, item in enumerate(profile.projects or []):
        items.append(
            {
                "id": f"project-{index}",
                "kind": "project",
                "label": str(item.get("name") or f"项目 {index + 1}"),
                "data": item,
                "text": _text(item),
            }
        )
    for index, value in enumerate(profile.certificates or []):
        items.append(
            {
                "id": f"additional-{index}",
                "kind": "additional",
                "label": value,
                "data": value,
                "text": value,
            }
        )
    return items


def analyze_job_description(
    job_description: str, profile: UserProfile
) -> dict[str, Any]:
    sources = _source_items(profile)
    matches: list[RequirementMatch] = []
    active_terms: list[str] = []

    for key, label, jd_terms, evidence_terms in CATEGORY_RULES:
        detected = [term for term in jd_terms if _contains(job_description, term)]
        if not detected:
            continue
        active_terms.extend(detected)
        evidence = [
            source["label"]
            for source in sources
            if any(_contains(source["text"], term) for term in evidence_terms)
        ]
        coverage = "strong" if len(evidence) >= 2 else "partial" if evidence else "gap"
        matches.append(
            RequirementMatch(
                key=key,
                label=label,
                coverage=coverage,
                evidence=evidence[:4],
            )
        )

    keywords = [
        keyword for keyword in TECH_KEYWORDS if _contains(job_description, keyword)
    ]
    return {
        "keywords": keywords,
        "requirements": [item.model_dump() for item in matches],
        "strong_count": sum(item.coverage == "strong" for item in matches),
        "partial_count": sum(item.coverage == "partial" for item in matches),
        "gap_count": sum(item.coverage == "gap" for item in matches),
        "active_terms": list(dict.fromkeys(active_terms)),
    }


def _score(text: str, analysis: dict[str, Any]) -> int:
    terms = [*analysis.get("keywords", []), *analysis.get("active_terms", [])]
    score = sum(3 for term in terms if _contains(text, term))
    for requirement in analysis.get("requirements", []):
        if any(label and _contains(text, label) for label in requirement["evidence"]):
            score += 5
    return score


def _entry_from_item(
    source_id: str, item: dict[str, Any], *, project: bool = False
) -> ResumeEntry:
    heading = str(item.get("name") if project else item.get("company") or "")
    subheading = str(item.get("role") if project else item.get("title") or "") or None
    if not project and item.get("department"):
        subheading = (
            f"{subheading} · {item['department']}" if subheading else item["department"]
        )
    bullets = [str(value) for value in (item.get("highlights") or []) if value]
    if not bullets and item.get("description"):
        bullets = [str(item["description"])]
    return ResumeEntry(
        source_id=source_id,
        heading=heading,
        subheading=subheading,
        period=_period(item),
        bullets=bullets[:3],
    )


def _build_summary(
    profile: UserProfile,
    job_title: str,
    ordered_skills: list[str],
) -> str:
    primary_education = (profile.education or [{}])[0]
    degree = profile.highest_degree or primary_education.get("degree")
    identity = "".join(
        value
        for value in [
            f"{profile.graduation_year}届" if profile.graduation_year else "",
            str(degree or ""),
        ]
        if value
    )
    sentences = [
        f"{identity or '应届候选人'}，目标岗位：{job_title}。",
    ]
    if ordered_skills:
        sentences.append(f"相关技能包括{'、'.join(ordered_skills[:6])}。")
    experience_types = []
    if profile.work_history:
        experience_types.append("实习")
    if profile.projects:
        experience_types.append("项目")
    if experience_types:
        sentences.append(f"具备{'与'.join(experience_types)}实践经历。")
    return "".join(sentences)


def build_tailored_content(
    job_title: str, job_description: str, profile: UserProfile
) -> tuple[dict[str, Any], TailoredResumeContent, dict[str, Any]]:
    analysis = analyze_job_description(job_description, profile)

    work = list(enumerate(profile.work_history or []))
    work.sort(key=lambda pair: _score(_text(pair[1]), analysis), reverse=True)
    experiences = [_entry_from_item(f"work-{index}", item) for index, item in work[:3]]

    work_corpus = _text(profile.work_history or [])
    projects = [
        (index, item)
        for index, item in enumerate(profile.projects or [])
        if not _contains(work_corpus, str(item.get("name") or ""))
    ]
    projects.sort(key=lambda pair: _score(_text(pair[1]), analysis), reverse=True)
    project_entries = [
        _entry_from_item(f"project-{index}", item, project=True)
        for index, item in projects[:3]
    ]

    skills = list(profile.skills or [])
    skills.sort(
        key=lambda skill: (
            0 if _contains(job_description, skill) else 1,
            (profile.skills or []).index(skill),
        )
    )

    education = [
        ResumeEducation(
            source_id=f"education-{index}",
            institution=str(item.get("institution") or ""),
            degree=str(item.get("degree") or ""),
            field_of_study=str(item.get("field_of_study") or ""),
            period=_period(item),
        )
        for index, item in enumerate(profile.education or [])
    ]

    additional = list(profile.certificates or [])
    additional.sort(key=lambda value: _score(value, analysis), reverse=True)
    if profile.languages:
        additional.append("语言：" + "、".join(profile.languages))

    content = TailoredResumeContent(
        target_role=job_title,
        summary=_build_summary(profile, job_title, skills),
        core_skills=skills[:18],
        experiences=experiences,
        projects=project_entries,
        education=education,
        additional=additional[:5],
    )
    validation = validate_tailored_content(profile, content, analysis)
    return analysis, content, validation


def validate_tailored_content(
    profile: UserProfile,
    content: TailoredResumeContent,
    analysis: dict[str, Any] | None = None,
) -> dict[str, Any]:
    valid_ids = {
        *(f"work-{i}" for i, _ in enumerate(profile.work_history or [])),
        *(f"project-{i}" for i, _ in enumerate(profile.projects or [])),
        *(f"education-{i}" for i, _ in enumerate(profile.education or [])),
    }
    errors: list[str] = []
    for entry in [*content.experiences, *content.projects, *content.education]:
        if entry.source_id not in valid_ids:
            errors.append(f"未知事实来源：{entry.source_id}")

    source_numbers = set(re.findall(r"\d+(?:\.\d+)?%?\+?", _text(profile.__dict__)))
    visible_content = content.model_dump()
    for section in ("experiences", "projects", "education"):
        for item in visible_content[section]:
            item.pop("source_id", None)
    output_numbers = set(re.findall(r"\d+(?:\.\d+)?%?\+?", _text(visible_content)))
    unsupported_numbers = sorted(output_numbers - source_numbers)
    if unsupported_numbers:
        errors.append("存在母版未记录的数字：" + "、".join(unsupported_numbers))

    performance_gap = any(
        item["key"] == "performance" and item["coverage"] != "strong"
        for item in (analysis or {}).get("requirements", [])
    )
    return {
        "valid": not errors,
        "errors": errors,
        "warnings": (
            ["性能与稳定性相关要求缺少直接履历证据"] if performance_gap else []
        ),
    }


def _set_run_font(run, size: float, *, bold: bool = False, color: str = "222222"):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _add_section_heading(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    _set_run_font(run, 11, bold=True, color="1F4E79")


def _add_entry(document: Document, entry: ResumeEntry) -> None:
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(13.5)
    table.columns[1].width = Cm(4.0)
    left, right = table.rows[0].cells
    heading = left.paragraphs[0]
    heading.paragraph_format.space_after = Pt(0)
    run = heading.add_run(entry.heading)
    _set_run_font(run, 9.5, bold=True)
    if entry.subheading:
        run = heading.add_run(f"  {entry.subheading}")
        _set_run_font(run, 9)
    period = right.paragraphs[0]
    period.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    period.paragraph_format.space_after = Pt(0)
    run = period.add_run(entry.period or "")
    _set_run_font(run, 8.5, color="555555")

    for bullet in entry.bullets:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.45)
        paragraph.paragraph_format.first_line_indent = Cm(-0.2)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(bullet)
        _set_run_font(run, 8.5)


def render_docx(
    profile: UserProfile,
    user: User,
    content: TailoredResumeContent,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9)

    name = "".join(part for part in [profile.last_name, profile.first_name] if part)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(1)
    run = title.add_run(name or "个人简历")
    _set_run_font(run, 16, bold=True, color="17365D")

    contact_values = [
        value
        for value in [
            profile.phone,
            profile.email,
            profile.portfolio_url,
            user.city,
        ]
        if value
    ]
    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    run = contact.add_run("  |  ".join(contact_values))
    _set_run_font(run, 8.5, color="555555")

    target = document.add_paragraph()
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target.paragraph_format.space_after = Pt(3)
    run = target.add_run(f"目标岗位：{content.target_role}")
    _set_run_font(run, 10, bold=True)

    _add_section_heading(document, "个人概述")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(content.summary)
    _set_run_font(run, 8.5)

    _add_section_heading(document, "核心技能")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(" · ".join(content.core_skills))
    _set_run_font(run, 8.5)

    _add_section_heading(document, "实习经历")
    for entry in content.experiences:
        _add_entry(document, entry)

    if content.projects:
        _add_section_heading(document, "项目经历")
        for entry in content.projects:
            _add_entry(document, entry)

    _add_section_heading(document, "教育经历")
    for item in content.education:
        entry = ResumeEntry(
            source_id=item.source_id,
            heading=item.institution,
            subheading=f"{item.field_of_study} · {item.degree}".strip(" ·"),
            period=item.period,
        )
        _add_entry(document, entry)

    if content.additional:
        _add_section_heading(document, "其他")
        for value in content.additional:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(1)
            run = paragraph.add_run(value)
            _set_run_font(run, 8.5)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def find_master_resume_filename(profile: UserProfile) -> str | None:
    for version in profile.resume_versions or []:
        if version.get("is_current"):
            return str(version.get("original_filename") or "") or None
    return profile.resume_original_filename
